#!/usr/bin/env python3
"""
共享底盘 v4.3 — Claude Code 代理核心
大猫和玄玄共用，永远走 Claude Code 代理模式。
读图走独立 ANTHROPIC_API_KEY，聊天永远走订阅。
# 底座已固定 v4.3 — 新功能只通过 trader/ 模块挂载，禁止修改此文件
"""
import os, json, logging, asyncio, time, base64, io, re, html as html_mod
from pathlib import Path
from dotenv import load_dotenv
from telegram import Update, BotCommand
from telegram.ext import Application, CommandHandler, MessageHandler, ContextTypes, filters
from telegram.constants import ParseMode
from telegram.error import RetryAfter


def _md_to_html(text: str) -> str:
    """Markdown → TG HTML（完整版）
    支持：代码块、行内代码、粗体、斜体、删除线、标题、
          分隔线、表格、有序/无序列表、链接、引用块
    """
    # ── 0. 保护区：先把代码块抽出来，防止内部被转换 ──
    _blocks = []
    def _save_block(m):
        _blocks.append(html_mod.escape(m.group(2)))
        return f"\x00CODEBLOCK{len(_blocks)-1}\x00"
    text = re.sub(r"```(\w*)\n(.*?)```", _save_block, text, flags=re.DOTALL)

    _inlines = []
    def _save_inline(m):
        _inlines.append(html_mod.escape(m.group(1)))
        return f"\x00CODEINLINE{len(_inlines)-1}\x00"
    text = re.sub(r"`([^`]+)`", _save_inline, text)

    # ── 1. 表格：| col | col | → 等宽文本 ──
    def _convert_table(m):
        lines = m.group(0).strip().split("\n")
        out = []
        for line in lines:
            stripped = line.strip()
            if re.match(r"^\|[\s\-:|]+\|$", stripped):
                continue  # 跳过分隔行 |---|---|
            out.append(stripped)
        return "<pre>" + "\n".join(out) + "</pre>"
    text = re.sub(r"(?:^\|.+\|$\n?){2,}", _convert_table, text, flags=re.MULTILINE)

    # ── 2. 标题 ##+ → 粗体（TG不支持原生标题） ──
    text = re.sub(r"^#{1,6}\s+(.+)$", r"<b>\1</b>", text, flags=re.MULTILINE)

    # ── 3. 分隔线 --- / *** / ___ → 空行 ──
    text = re.sub(r"^[\s]*[-*_]{3,}[\s]*$", "", text, flags=re.MULTILINE)

    # ── 4. 引用块 > text → 竖线缩进 ──
    text = re.sub(r"^>\s?(.*)$", r"┃ \1", text, flags=re.MULTILINE)

    # ── 5. 粗斜体 ***text*** → 粗+斜 ──
    text = re.sub(r"\*\*\*(.+?)\*\*\*", r"<b><i>\1</i></b>", text)

    # ── 6. 粗体 **text** → <b> ──
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)

    # ── 7. 斜体 *text* → <i>（不匹配emoji旁的孤立*） ──
    text = re.sub(r"(?<!\w)\*([^\s*](?:.*?[^\s*])?)\*(?!\w)", r"<i>\1</i>", text)

    # ── 8. 删除线 ~~text~~ → <s> ──
    text = re.sub(r"~~(.+?)~~", r"<s>\1</s>", text)

    # ── 9. 链接 [text](url) → <a> ──
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r'<a href="\2">\1</a>', text)

    # ── 10. 有序列表美化：数字. → 数字. （保持原样，TG纯文本就行） ──
    # ── 11. 无序列表：- / * 开头 → • ──
    text = re.sub(r"^[\s]*[-*]\s+", "• ", text, flags=re.MULTILINE)

    # ── 12. 恢复代码块 ──
    for i, code in enumerate(_blocks):
        text = text.replace(f"\x00CODEBLOCK{i}\x00", f"<pre><code>{code}</code></pre>")
    for i, code in enumerate(_inlines):
        text = text.replace(f"\x00CODEINLINE{i}\x00", f"<code>{code}</code>")

    # ── 13. HTML 特殊字符转义（v2 2026-04-26：全裸字符转义+合法标签保护）──
    # 旧版只转 &，导致 LLM 输出 "score<60" / "if x<5" / "a&b" 时 TG 解析失败
    # 触发降级路径剥光所有标签 → 复制按钮+加粗+代码块全消失（"美化反复消失"根因）
    # 13.1 先把上面已生成的合法 HTML 标签抽出来占位
    _legal = []
    def _save_tag(m):
        _legal.append(m.group(0))
        return f"\x00TAG{len(_legal)-1}\x00"
    text = re.sub(
        r'</?(?:b|i|s|u|code|pre|a|tg-spoiler)\b[^>]*>',
        _save_tag, text
    )
    # 13.2 此时 text 里所有 < > & 都是裸字符，全部转义
    text = re.sub(r"&(?!amp;|lt;|gt;|quot;|#\d+;)", "&amp;", text)
    text = text.replace('<', '&lt;').replace('>', '&gt;')
    # 13.3 恢复合法标签
    for i, t in enumerate(_legal):
        text = text.replace(f"\x00TAG{i}\x00", t)

    return text

MODELS = [
    "claude-sonnet-4-6",
    "claude-opus-4-6",
    "claude-opus-4-7",
    "claude-haiku-4-5-20251001",
]

def _state_file(bot_dir):
    p = Path(f"/root/{bot_dir}/data/mode.json")
    p.parent.mkdir(parents=True, exist_ok=True)
    return p

def load_state(bot_dir):
    try:
        s = json.loads(_state_file(bot_dir).read_text())
        if "model" not in s:
            s["model"] = MODELS[0]
        return s
    except Exception:
        return {"model": MODELS[0]}

def save_state(bot_dir, state):
    _state_file(bot_dir).write_text(json.dumps(state, ensure_ascii=False))

def model_short(model):
    return model.replace("claude-", "").replace("-20251001", "")


# ── Claude Code 完整 Agent（聊天永远走这里）──
async def claudecode_gen(prompt, add_dir="/root", bot_dir="damao", model=None):
    async def _gen(step_info):
        import json as _json
        TOOL_ICONS = {
            "Bash": "🖥️", "Read": "📘", "Write": "📒", "Edit": "📙",
            "Grep": "🔍", "Glob": "🗂️", "WebSearch": "🌐", "WebFetch": "🌏",
            "Task": "🦾", "TodoWrite": "📌", "Skill": "✨",
        }
        step_info["text"] = "启动中"
        step_info["actions"] = []
        cmd = ["claude", "--print", "--verbose",
               "--output-format", "stream-json",
               "--permission-mode", "auto",
               "--continue",
               "--add-dir", add_dir]
        if model:
            cmd.extend(["--model", model])
        cmd.extend(["-p", prompt])
        proc = await asyncio.create_subprocess_exec(
            *cmd, stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE, cwd=add_dir,
            env={k: v for k, v in os.environ.items() if k != "ANTHROPIC_API_KEY"},
            limit=1024 * 1024 * 32)
        final_result = ""
        async def _read():
            nonlocal final_result
            async for raw_line in proc.stdout:
                line = raw_line.decode("utf-8").strip()
                if not line:
                    continue
                try:
                    obj = _json.loads(line)
                except Exception:
                    continue
                typ = obj.get("type")
                if typ == "assistant":
                    for item in obj.get("message", {}).get("content", []):
                        if item.get("type") == "tool_use":
                            tool_name = item.get("name", "")
                            icon = TOOL_ICONS.get(tool_name, "⚙️")
                            inp = item.get("input", {})
                            brief = str(inp.get("command") or inp.get("file_path") or
                                        inp.get("pattern") or inp.get("query") or
                                        inp.get("prompt") or "")[:50]
                            step_info["actions"].append(f"{icon} {brief or tool_name}")
                            step_info["text"] = f"{icon} {tool_name}"
                elif typ == "result":
                    final_result = obj.get("result", "").strip()
        try:
            await asyncio.wait_for(_read(), timeout=600)
        except asyncio.TimeoutError:
            try:
                proc.kill()
            except Exception:
                pass
            step_info["text"] = "超时"
            yield "⏱️ Claude响应超时（600s），已终止进程"
            return
        _, stderr_bytes = await proc.communicate()
        step_info["text"] = "完成"
        yield final_result or stderr_bytes.decode("utf-8")[:300] or "（无输出）"
    return _gen


# ── API 读图（仅用于图片，不用于聊天）──
async def api_gen_image(prompt, api_key, model, system_prompt, image_b64):
    async def _gen(step_info):
        import anthropic
        client = anthropic.Anthropic(api_key=api_key)
        content = [
            {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": image_b64}},
            {"type": "text", "text": prompt or "请描述这张图片"}
        ]
        step_info["text"] = "读图中"
        loop = asyncio.get_event_loop()
        queue = asyncio.Queue()
        def _stream():
            try:
                with client.messages.stream(model=model, max_tokens=4096, system=system_prompt,
                    messages=[{"role": "user", "content": content}]) as stream:
                    step_info["text"] = "生成描述"
                    for text in stream.text_stream:
                        loop.call_soon_threadsafe(queue.put_nowait, text)
                loop.call_soon_threadsafe(queue.put_nowait, None)
            except Exception as e:
                loop.call_soon_threadsafe(queue.put_nowait, f"\n⚠️ 读图失败: {e}")
                loop.call_soon_threadsafe(queue.put_nowait, None)
        loop.run_in_executor(None, _stream)
        while True:
            chunk = await queue.get()
            if chunk is None:
                break
            yield chunk
    return _gen


# ── 计时器 UI ──
async def ask_with_timer(update, gen_coro, model, voice_reply_fn=None):
    thinking_msg = await update.message.reply_text(f"🐦 乌鸦团队 · 启动中... 0s")
    start_time = time.time()
    step_info = {"text": "启动中"}
    async def ticker():
        # 基础节奏 3s（多 ticker 并发也不易触 TG 单 chat 限流）；
        # RetryAfter 命中后用其 retry_after，其它瞬时错误指数退避封顶 30s；成功即衰减回基础节奏。
        BASE_INTERVAL = 3.0
        MAX_BACKOFF = 30.0
        backoff = BASE_INTERVAL
        last_text = None
        while True:
            await asyncio.sleep(backoff)
            elapsed = int(time.time() - start_time)
            actions = step_info.get("actions")
            if actions:
                lines = "\n".join(f"  {a}" for a in actions[-10:])
                txt = f"⚡ {elapsed}s\n{lines}"
            else:
                txt = f"🐦 乌鸦团队 · {step_info['text']}... {elapsed}s"
            if txt == last_text:
                continue   # 文本未变就不打 API（TG 也会回 BadRequest "not modified"）
            try:
                await thinking_msg.edit_text(txt)
                last_text = txt
                backoff = BASE_INTERVAL
            except RetryAfter as e:
                backoff = min(max(float(getattr(e, "retry_after", 5)), BASE_INTERVAL), MAX_BACKOFF)
            except Exception:
                backoff = min(backoff * 2, MAX_BACKOFF)
    ticker_task = asyncio.create_task(ticker())
    full_text = ""
    try:
        async for chunk in gen_coro(step_info):
            full_text += chunk
    except Exception as e:
        full_text = f"⚠️ 错误: {e}"
    finally:
        ticker_task.cancel()
    try:
        await thinking_msg.delete()
    except Exception:
        pass
    if not full_text.strip():
        full_text = "（无输出）"
    # 2026-04-27 新增：voice_reply_fn 用于语音输入场景，让回复也走 TTS
    if voice_reply_fn:
        try:
            await voice_reply_fn(update, full_text)
            return
        except Exception as e:
            logging.getLogger("core").warning(f"voice_reply 失败 fallback 文字: {e}")
    html_text = _md_to_html(full_text)
    chunks = [html_text[i:i+4000] for i in range(0, len(html_text), 4000)]
    for chunk in chunks:
        try:
            await update.message.reply_text(chunk, parse_mode=ParseMode.HTML)
        except Exception as e:
            logging.getLogger("core").warning(f"HTML解析失败，降级纯文本: {e}")
            # 剥掉HTML标签，保留可读文本
            plain = re.sub(r"<[^>]+>", "", chunk)
            await update.message.reply_text(plain)


def create_and_run_bot(env_path, claude_add_dir=None):
    load_dotenv(env_path)
    bot_token     = os.environ["BOT_TOKEN"]
    admin_id      = int(os.environ["ADMIN_ID"])
    anthropic_key = os.getenv("ANTHROPIC_API_KEY", "")
    bot_name      = os.getenv("BOT_NAME", "Bot")
    bot_dir       = os.getenv("BOT_DIR", "bot")
    system_prompt = os.getenv("SYSTEM_PROMPT", "你是AI助手。")
    log_dir = Path(f"/root/{bot_dir}/logs")
    log_dir.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(format="%(asctime)s [%(levelname)s] %(message)s", level=logging.INFO,
        handlers=[logging.StreamHandler(), logging.FileHandler(log_dir / "bot.log", encoding="utf-8")])
    logger = logging.getLogger(bot_name)
    if bot_dir == "damao":
        _add_dir = claude_add_dir or "/root"
    elif bot_dir == "maomao":
        _add_dir = "/root/maomao"
    elif bot_dir == "tiantian":
        _add_dir = "/root/tiantian"
    else:
        _add_dir = claude_add_dir or "/root"

    def admin_only(func):
        async def wrapper(update, ctx, *args, **kwargs):
            uid = update.effective_user.id
            if uid != admin_id:
                uname = update.effective_user.username or ""
                txt = update.message.text if update.message else ""
                logger.warning(f"[admin_only] 拒绝: user_id={uid} username=@{uname} text={txt!r} expected_admin={admin_id}")
                return
            return await func(update, ctx, *args, **kwargs)
        return wrapper

    async def ask_claude(update, prompt, image_b64=None):
        from datetime import datetime as _dt
        now_str = _dt.now().strftime("%Y-%m-%d %H:%M:%S")
        prompt = f"[当前时间: {now_str}] {prompt}"
        state = load_state(bot_dir)
        model = state.get("model", MODELS[0])
        if image_b64:
            if not anthropic_key:
                await update.message.reply_text("⚠️ ANTHROPIC_API_KEY 未设置，无法读图")
                return
            # 2026-04-27 老大指令: 读图强制 Haiku 4.5（省 96% 费用），不跟 mode.json 走
            # 原因: api_gen_image 走 ANTHROPIC_API_KEY 按量计费，读图不需要 Opus 智力
            gen = await api_gen_image(prompt, anthropic_key, "claude-haiku-4-5-20251001", system_prompt, image_b64)
        else:
            gen = await claudecode_gen(prompt, add_dir=_add_dir, bot_dir=bot_dir, model=model)
        await ask_with_timer(update, gen, model)

    @admin_only
    async def cmd_start(update, ctx):
        state = load_state(bot_dir)
        icon = '🐱' if bot_name == '大猫' else '🐦' if bot_name == '玄玄' else '🐶' if bot_name == '贝贝' else '🐾'
        await update.message.reply_text(
            f"{icon} <b>{bot_name} v4.3</b>\n\n"
            f"模型: <code>{state['model']}</code>\n\n"
            f"/model — 切换模型\n/help — 全部命令",
            parse_mode=ParseMode.HTML)

    @admin_only
    async def cmd_help(update, ctx):
        state = load_state(bot_dir)
        if bot_dir == "tiantian":
            await update.message.reply_text(
                f"<b>{bot_name}命令</b>\n\n"
                f"模型: <code>{model_short(state['model'])}</code>\n\n"
                f"/model   — 切换模型\n"
                f"/status  — Bot状态\n"
                f"/restart — 重启\n/stop — 关闭\n\n"
                f"/1 — 查持仓\n/2 — 查余额\n"
                f"/3 — 现货→合约\n/4 — 合约→现货\n"
                f"/5 — 现货→资金\n/6 — 资金→现货\n\n"
                f"<i>发文字 = 问{bot_name}\n发图片 = Vision读图</i>",
                parse_mode=ParseMode.HTML)
            return
        await update.message.reply_text(
            f"<b>{bot_name}命令</b>\n\n"
            f"模型: <code>{model_short(state['model'])}</code>\n\n"
            f"/model  — 循环切换模型\n"
            f"/mode   — 查看状态\n/status — 四Bot服务状态\n"
            f"/ping   — 心跳\n/log    — 查看日志\n\n"
            f"<i>发文字 = 问{bot_name}\n发图片 = Vision读图（消耗API额度）</i>",
            parse_mode=ParseMode.HTML)

    @admin_only
    async def cmd_model(update, ctx):
        state = load_state(bot_dir)
        cur = state.get("model", MODELS[0])
        try:
            idx = MODELS.index(cur)
        except ValueError:
            idx = 0
        nxt = MODELS[(idx + 1) % len(MODELS)]
        state["model"] = nxt
        save_state(bot_dir, state)
        await update.message.reply_text(
            f"✅ 模型已切换\n\n`{model_short(cur)}`\n↓\n`{model_short(nxt)}`",
            parse_mode=ParseMode.MARKDOWN)

    @admin_only
    async def cmd_mode(update, ctx):
        if bot_dir == "tiantian":
            state = load_state(bot_dir)
            await update.message.reply_text(
                f"<b>{bot_name} 状态</b>\n\n模型: <code>{state['model']}</code>",
                parse_mode=ParseMode.HTML)
            return
        import shutil
        state = load_state(bot_dir)
        api_ok = "✅" if anthropic_key else "❌"
        cli_ok = "✅" if shutil.which("claude") else "⚠️ 未找到"
        await update.message.reply_text(
            f"<b>{bot_name} 状态</b>\n\n"
            f"模型: <code>{state['model']}</code>\n\n"
            f"📦 Claude Code: {cli_ok}\n"
            f"🖼 读图API: {api_ok}",
            parse_mode=ParseMode.HTML)

    @admin_only
    async def cmd_ping(update, ctx):
        state = load_state(bot_dir)
        await update.message.reply_text(
            f"🏓 {bot_name}存活 · `{model_short(state['model'])}`",
            parse_mode=ParseMode.MARKDOWN)

    @admin_only
    async def cmd_status(update, ctx):
        import subprocess as sp
        if bot_dir == "tiantian":
            svcs = ["tiantian"]
        else:
            svcs = ["damao", "maomao", "tiantian", "baobao"]
        lines = []
        for svc in svcs:
            r = sp.run(["systemctl", "is-active", svc], capture_output=True, text=True)
            icon = "✅" if r.stdout.strip() == "active" else "❌"
            lines.append(f"{icon} {svc}: {r.stdout.strip()}")
        await update.message.reply_text(
            "<b>服务状态</b>\n\n" + "\n".join(lines),
            parse_mode=ParseMode.HTML)

    @admin_only
    async def cmd_log(update, ctx):
        if bot_dir == "tiantian":
            await update.message.reply_text("⛔ 此命令未开放")
            return
        try:
            lines = (log_dir / "bot.log").read_text().splitlines()[-50:]
            text = "\n".join(lines) or "（日志为空）"
            await update.message.reply_text(f"<pre>{text[-3500:]}</pre>", parse_mode=ParseMode.HTML)
        except Exception as e:
            await update.message.reply_text(f"读取日志失败: {e}")

    _CONFIRM_WORDS = {"确认", "ok", "OK", "好", "是", "确定", "yes", "y", "Y"}
    _CANCEL_WORDS  = {"取消", "cancel", "不", "算了", "不要", "no", "n", "N"}

    @admin_only
    async def handle_text(update, ctx, text_override=None):
        if not update.message or not update.message.text:
            return
        user_text = text_override if text_override is not None else update.message.text
        if bot_dir in ("maomao", "tiantian"):
            import sys
            if '/root/maomao' not in sys.path:
                sys.path.insert(0, '/root/maomao')
            # 2026-04-21 入口路由统一切到 trader.multi.dispatch
            # （替代老 trader.router/order/exchange 单账户路径，
            #  根因：老路径默认币安1+全仓，导致跨账户指令静默落点错误）
            role = {"maomao": "玄玄", "tiantian": "天天"}.get(bot_dir, bot_dir)
            try:
                from trader.multi.dispatch import try_dispatch
                reply, status = try_dispatch(role, user_text)
            except Exception as e:
                logger.error(f"[dispatch] {role} 异常: {e}", exc_info=True)
                reply, status = (f"❌ 派发异常: {e}", "err")
            if status != "none":
                try:
                    await update.message.reply_text(reply, parse_mode=ParseMode.HTML)
                except Exception:
                    await update.message.reply_text(reply)
                try:
                    from trader.trade_log import log_trade
                    log_trade(raw_text=user_text,
                              result=reply if status == "ok" else None,
                              error=reply if status == "err" else None)
                except ImportError:
                    pass
                return
        await ask_claude(update, user_text)

    async def handle_callback(update, ctx):
        await update.callback_query.answer()

    @admin_only
    async def handle_photo(update, ctx):
        photo = update.message.photo[-1]
        file = await ctx.bot.get_file(photo.file_id)
        buf = io.BytesIO()
        await file.download_to_memory(buf)
        buf.seek(0)
        image_b64 = base64.b64encode(buf.read()).decode()
        caption = update.message.caption or "请描述这张图片"
        await ask_claude(update, caption, image_b64)

    @admin_only
    async def handle_document(update, ctx):
        # 2026-04-27 老大要求：手机也能传文件
        # 1) 一律存到 /root/shared/inbox/<时间戳>_<文件名> 留底
        # 2) 按类型解析（PDF / docx / xlsx / 文本）→ 走 ask_claude 订阅，不消耗 API key
        # 3) 解析失败/不支持的类型也通知 Claude 已存档路径，让大猫主动 Read
        # 4) 50 MB 是 Bot API 硬限制，超过会被 TG 直接拒绝
        from datetime import datetime as _dt
        doc = update.message.document
        fname = doc.file_name or "file"
        mime = doc.mime_type or ""
        size = doc.file_size or 0
        file = await ctx.bot.get_file(doc.file_id)
        buf = io.BytesIO()
        await file.download_to_memory(buf)
        buf.seek(0)
        raw = buf.getvalue()
        caption = update.message.caption or ""

        # 一律存盘留底
        inbox = Path("/root/shared/inbox")
        inbox.mkdir(parents=True, exist_ok=True)
        ts = _dt.now().strftime("%Y%m%d-%H%M%S")
        save_path = inbox / f"{ts}_{fname}"
        save_path.write_bytes(raw)
        archive_note = f"\n\n（已存档：{save_path}，大小 {size}B）"

        MAX_TEXT = 50000  # Claude context window 充足，从 12000 放大到 50000

        try:
            if mime == "application/pdf" or fname.lower().endswith(".pdf"):
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(raw))
                text = "\n".join(p.extract_text() or "" for p in reader.pages)
                prompt = f"📄 PDF 文件：{fname}\n{caption}\n\n内容：\n\n{text[:MAX_TEXT]}{archive_note}"
                await ask_claude(update, prompt)

            elif fname.lower().endswith(".docx"):
                from docx import Document as _DocxDoc
                d = _DocxDoc(io.BytesIO(raw))
                text = "\n".join(p.text for p in d.paragraphs if p.text)
                # 表格内容
                for tbl in d.tables:
                    for row in tbl.rows:
                        text += "\n| " + " | ".join(c.text.strip() for c in row.cells) + " |"
                prompt = f"📝 Word 文档：{fname}\n{caption}\n\n内容：\n\n{text[:MAX_TEXT]}{archive_note}"
                await ask_claude(update, prompt)

            elif fname.lower().endswith((".xlsx", ".xls")):
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
                text_lines = []
                for ws in wb.worksheets:
                    text_lines.append(f"━━ Sheet: {ws.title} ━━")
                    for row in ws.iter_rows(values_only=True, max_rows=200):
                        text_lines.append("\t".join(str(c) if c is not None else "" for c in row))
                text = "\n".join(text_lines)
                prompt = f"📊 Excel：{fname}\n{caption}\n\n内容：\n\n{text[:MAX_TEXT]}{archive_note}"
                await ask_claude(update, prompt)

            elif (mime.startswith("text/")
                  or fname.lower().endswith((".txt", ".md", ".yaml", ".yml", ".json",
                                             ".py", ".js", ".ts", ".sh", ".csv",
                                             ".log", ".html", ".xml", ".sql", ".toml", ".ini"))):
                text = raw.decode("utf-8", errors="replace")
                prompt = f"📃 文本文件：{fname}\n{caption}\n\n内容：\n\n{text[:MAX_TEXT]}{archive_note}"
                await ask_claude(update, prompt)

            else:
                # 不支持解析的类型（zip/exe/bin 等）：通知 Claude 已存档让大猫主动 Read
                prompt = (f"📦 收到文件 {fname}（{mime or '未知类型'}, {size}B），"
                          f"系统不支持自动解析，已存档于 {save_path}。\n"
                          f"用户备注：{caption or '无'}\n"
                          f"如需查看请用 Read 工具读取该路径，或建议用户用其他方式发送内容。")
                await ask_claude(update, prompt)

        except Exception as e:
            # 解析失败 fallback：通知 Claude 文件存了，但解析挂了
            prompt = (f"⚠️ 文件 {fname} 解析失败（{type(e).__name__}: {e}），"
                      f"已存档于 {save_path}。请考虑直接 Read 该文件或建议用户换格式。")
            await ask_claude(update, prompt)

    # 2026-04-27 视频处理：ffmpeg 截帧 + Haiku Vision 看图描述 + 总结
    @admin_only
    async def handle_video(update, ctx):
        from datetime import datetime as _dt
        import subprocess, tempfile, os as _os
        msg = update.message
        v = msg.video or msg.video_note
        if not v:
            return
        fname = getattr(v, "file_name", None) or f"video_{v.file_id[:10]}.mp4"
        size = v.file_size or 0
        duration = getattr(v, "duration", 0) or 0
        caption = msg.caption or ""

        await msg.reply_text(f"🎬 收到视频 {fname}（{size//1024}KB / {duration}s），开始截帧分析...")

        file = await ctx.bot.get_file(v.file_id)
        buf = io.BytesIO()
        await file.download_to_memory(buf)
        raw = buf.getvalue()

        # 存盘留底
        inbox = Path("/root/shared/inbox")
        inbox.mkdir(parents=True, exist_ok=True)
        ts = _dt.now().strftime("%Y%m%d-%H%M%S")
        save_path = inbox / f"{ts}_{fname}"
        save_path.write_bytes(raw)

        # ffmpeg 截帧：每 30s 一帧，最多 5 帧（防长视频炸钱）
        max_frames = 5
        interval = max(30, duration // max_frames) if duration else 30
        timestamps = [i * interval for i in range(max_frames) if not duration or i * interval < duration]
        if not timestamps:
            timestamps = [0]

        frame_descriptions = []
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                for idx, ts_sec in enumerate(timestamps):
                    out = _os.path.join(tmpdir, f"frame_{idx}.jpg")
                    cmd = ["ffmpeg", "-y", "-ss", str(ts_sec), "-i", str(save_path),
                           "-frames:v", "1", "-q:v", "3", out]
                    r = subprocess.run(cmd, capture_output=True, timeout=20)
                    if r.returncode != 0 or not _os.path.exists(out):
                        continue
                    with open(out, "rb") as f:
                        b64 = base64.b64encode(f.read()).decode()
                    # 调 Haiku 描述这一帧
                    desc_chunks = []
                    gen = await api_gen_image(
                        f"用 1-2 句话描述这帧画面（视频 {fname} 第 {ts_sec}s）",
                        anthropic_key, "claude-haiku-4-5-20251001", system_prompt, b64
                    )
                    async for chunk in gen({"text": "", "actions": []}):
                        desc_chunks.append(chunk)
                    frame_descriptions.append(f"[{ts_sec}s] {''.join(desc_chunks)}")
        except Exception as e:
            frame_descriptions.append(f"⚠️ 截帧失败: {e}")

        # 拼成 prompt 给 Claude（订阅）总结
        frames_text = "\n".join(frame_descriptions) if frame_descriptions else "（无帧）"
        prompt = (f"🎬 视频文件：{fname}（{duration}s, {size//1024}KB）\n"
                  f"用户备注：{caption or '无'}\n\n"
                  f"截了 {len(frame_descriptions)} 帧，每帧 Haiku 描述：\n{frames_text}\n\n"
                  f"（视频已存档：{save_path}）\n"
                  f"请基于这些帧描述，给个整体总结。")
        await ask_claude(update, prompt)

    # 读 voice profile（control.yaml -> voices.{bot_dir}）
    def _voice_profile():
        try:
            import yaml as _y
            ctl = _y.safe_load(open("/root/maomao/control.yaml").read()) or {}
            return (ctl.get("voices") or {}).get(bot_dir, {}) or {}
        except Exception:
            return {}

    # TTS 三引擎：volc（豆包大模型，最佳中文 + 真童音）/ edge / openai
    async def _tts_bytes(text: str, engine: str, voice: str, instructions: str = "") -> bytes:
        text = (text or "")[:4000]
        if engine == "volc":
            # 火山豆包大模型 TTS V3 unidirectional（NDJSON 流式响应）
            import requests as _req, base64 as _b64, json as _json
            headers = {
                "X-Api-Key": os.getenv("VOLC_TTS_API_KEY", ""),
                "X-Api-App-Id": os.getenv("VOLC_APP_ID", ""),
                "X-Api-Resource-Id": "volc.service_type.10029",
                "Content-Type": "application/json",
            }
            body = {
                "user": {"uid": "bot"},
                "req_params": {
                    "text": text,
                    "speaker": voice,
                    "audio_params": {"format": "mp3", "sample_rate": 24000},
                }
            }
            loop = asyncio.get_event_loop()
            resp = await loop.run_in_executor(
                None, lambda: _req.post(
                    "https://openspeech.bytedance.com/api/v3/tts/unidirectional",
                    json=body, headers=headers, timeout=30))
            chunks = []
            for line in resp.text.split("\n"):
                if not line.strip():
                    continue
                try:
                    obj = _json.loads(line)
                    if obj.get("code") == 0 and obj.get("data"):
                        chunks.append(_b64.b64decode(obj["data"]))
                except Exception:
                    pass
            return b"".join(chunks)
        if engine == "edge":
            import edge_tts
            comm = edge_tts.Communicate(text, voice=voice)
            chunks = []
            async for chunk in comm.stream():
                if chunk.get("type") == "audio":
                    chunks.append(chunk["data"])
            return b"".join(chunks)
        # 默认 openai
        import openai as _oai
        _client = _oai.OpenAI(api_key=os.getenv("OPENAI_API_KEY", ""))
        kwargs = {"model": "gpt-4o-mini-tts", "voice": voice, "input": text}
        if instructions:
            kwargs["instructions"] = instructions
        loop = asyncio.get_event_loop()
        resp = await loop.run_in_executor(None, lambda: _client.audio.speech.create(**kwargs))
        return resp.content

    async def _send_voice_reply(update, text: str):
        """按 bot_dir 选引擎 + voice，转成 mp3 发 voice 给用户"""
        prof = _voice_profile()
        if not prof.get("tts_enabled"):
            await update.message.reply_text(text)
            return
        engine = prof.get("engine", "edge")
        voice = prof.get("voice", "zh-CN-XiaoxiaoNeural")
        instructions = prof.get("instructions", "")
        try:
            mp3 = await _tts_bytes(text, engine=engine, voice=voice, instructions=instructions)
            await update.message.reply_voice(io.BytesIO(mp3), caption=text[:1024] if len(text) > 200 else None)
        except Exception as e:
            logger.warning(f"[TTS] 失败 fallback 文字: {e}")
            await update.message.reply_text(text)

    # 2026-04-27 音频处理：OpenAI Whisper STT → 走 dispatch 或 ask_claude → TTS 回复
    # 用途：玄玄/天天端语音下单/查询/聊天，省手打字时间
    # 闭环：你发 voice → Whisper 转字 → 处理 → TTS 转字 → 你听 voice
    # 缺 OPENAI_API_KEY 时优雅 fallback：只存盘 + 通知大猫
    @admin_only
    async def handle_audio(update, ctx):
        from datetime import datetime as _dt
        msg = update.message
        a = msg.voice or msg.audio
        if not a:
            return
        fname = getattr(a, "file_name", None) or f"audio_{a.file_id[:10]}.ogg"
        size = a.file_size or 0
        duration = getattr(a, "duration", 0) or 0
        mime = getattr(a, "mime_type", "audio/ogg") or "audio/ogg"
        caption = msg.caption or ""

        file = await ctx.bot.get_file(a.file_id)
        buf = io.BytesIO()
        await file.download_to_memory(buf)
        raw = buf.getvalue()

        # 一律存盘
        inbox = Path("/root/shared/inbox")
        inbox.mkdir(parents=True, exist_ok=True)
        ts = _dt.now().strftime("%Y%m%d-%H%M%S")
        save_path = inbox / f"{ts}_{fname}"
        save_path.write_bytes(raw)

        # 尝试 Whisper 转文字
        openai_key = os.getenv("OPENAI_API_KEY", "")
        if not openai_key:
            # Fallback：缺 key 走存档 + 通知
            prompt = (f"🎤 收到音频 {fname}（{mime}, {duration}s, {size//1024}KB）\n"
                      f"用户备注：{caption or '无'}\n"
                      f"已存档：{save_path}\n\n"
                      f"⚠️ OPENAI_API_KEY 未配置，无法转文字。请老大 SSH 写入 master 后重启。")
            await ask_claude(update, prompt)
            return

        # 2026-04-27 删冗余提示：直接转文字，不发"收到音频"中间消息
        try:
            import openai
            client = openai.OpenAI(api_key=openai_key)
            with open(save_path, "rb") as f:
                resp = client.audio.transcriptions.create(
                    model="whisper-1",
                    file=f,
                    language="zh",
                    response_format="text",
                )
            transcript = (resp if isinstance(resp, str) else getattr(resp, "text", str(resp))).strip()
        except Exception as e:
            await msg.reply_text(f"⚠️ Whisper 转文字失败: {e}")
            await ask_claude(update, f"🎤 音频已存档于 {save_path}，但 Whisper 转文字失败：{e}")
            return

        if not transcript:
            await msg.reply_text("⚠️ 转文字结果为空，可能音频太短或无人声")
            return
        # 2026-04-27 老大要求：删除"📝 听到：xxx"中间提示，对话框只保留最终回复

        # 先尝试 dispatch（交易指令命中即直接执行，不进 Claude）
        # 仅 maomao / tiantian 走 dispatch，damao / baobao 直接 ask_claude
        dispatched = False
        if bot_dir in ("maomao", "tiantian"):
            try:
                import sys as _sys
                if "/root/maomao" not in _sys.path:
                    _sys.path.insert(0, "/root/maomao")
                from trader.multi.dispatch import try_dispatch
                role = {"maomao": "玄玄", "tiantian": "天天"}[bot_dir]
                full_text = f"{caption} {transcript}".strip() if caption else transcript
                disp_result = try_dispatch(role, full_text)
                # try_dispatch 返回 tuple (reply: str|None, source: str)
                reply, _src = disp_result if isinstance(disp_result, tuple) else (disp_result, "")
                if reply:
                    # dispatch 命中 → 用 TTS 语音回
                    await _send_voice_reply(update, reply)
                    dispatched = True
            except Exception as e:
                logger.warning(f"[handle_audio] dispatch 失败: {e}")

        # dispatch 没命中 → 走 Claude 自然对话 + 语音回（闭环）
        if not dispatched:
            full_prompt = f"{caption}\n\n（语音输入转文字）{transcript}" if caption else f"（语音输入转文字）{transcript}"
            # 手动跑 claudecode_gen + ask_with_timer 传 voice_reply_fn，让回复走 TTS
            state = load_state(bot_dir)
            model = state.get("model", MODELS[0])
            gen = await claudecode_gen(full_prompt, add_dir=_add_dir, bot_dir=bot_dir, model=model)
            # voice_reply_fn 决定是否 TTS：profile.tts_enabled=False 时退化为文字回
            await ask_with_timer(update, gen, model, voice_reply_fn=_send_voice_reply)

    @admin_only
    async def cmd_restart(update, ctx):
        await update.message.reply_text(f"🔄 {bot_name} 重启中...")
        import subprocess
        subprocess.Popen(["systemctl", "restart", bot_dir])

    @admin_only
    async def cmd_stop_bot(update, ctx):
        await update.message.reply_text(f"🛑 {bot_name} 关闭中...")
        import subprocess
        subprocess.Popen(["systemctl", "stop", bot_dir])

    async def error_handler(update, ctx):
        logger.error("Exception:", exc_info=ctx.error)
        _bot_log("error", str(ctx.error)[:200])

    # ── 多账户快捷卡片（2026-04-19 拆分：持仓/余额独立） ──
    _ACC_ROLE = {"maomao": "玄玄", "tiantian": "天天", "damao": "大猫"}.get(bot_dir)

    def _mk_card_handler(kind, acc_name):
        """kind ∈ {'pos','bal'}，返回对应 account 的持仓卡或余额卡"""
        @admin_only
        async def _h(update, ctx):
            try:
                if kind == "pos":
                    from shared.query_cards import render_positions_card
                    card = render_positions_card(_ACC_ROLE, acc_name)
                else:
                    from shared.query_cards import render_wallet_card
                    card = render_wallet_card(_ACC_ROLE, acc_name)
                await update.message.reply_text(card, parse_mode=ParseMode.HTML)
            except PermissionError as e:
                await update.message.reply_text(f"⛔ {e}")
            except Exception as e:
                logger.error(f"{kind} {acc_name} 失败: {e}", exc_info=True)
                await update.message.reply_text(f"❌ {acc_name} 查询失败: {e}")
        return _h

    @admin_only
    async def cmd_all_card(update, ctx):
        try:
            from shared.query_cards import render_all_card
            card = render_all_card(_ACC_ROLE)
            await update.message.reply_text(card, parse_mode=ParseMode.HTML)
        except PermissionError as e:
            await update.message.reply_text(f"⛔ {e}")
        except Exception as e:
            logger.error(f"全查失败: {e}", exc_info=True)
            await update.message.reply_text(f"❌ 全查失败: {e}")

    cmd_pos1 = _mk_card_handler("pos", "币安1")
    cmd_pos2 = _mk_card_handler("pos", "币安2")
    cmd_pos3 = _mk_card_handler("pos", "币安3")
    cmd_pos4 = _mk_card_handler("pos", "币安4")
    cmd_bal1 = _mk_card_handler("bal", "币安1")
    cmd_bal2 = _mk_card_handler("bal", "币安2")
    cmd_bal3 = _mk_card_handler("bal", "币安3")
    cmd_bal4 = _mk_card_handler("bal", "币安4")

    @admin_only
    async def cmd_q_positions(update, ctx):
        """/1 全账户持仓（纯持仓，按 role 过滤；迁到 trader.multi 四账户）"""
        try:
            from shared.query_cards import render_all_positions_card
            card = render_all_positions_card(_ACC_ROLE)
            if len(card) > 4000:
                card = card[:4000] + "\n..."
            await update.message.reply_text(card, parse_mode=ParseMode.HTML)
        except PermissionError as e:
            await update.message.reply_text(f"⛔ {e}")
        except Exception as e:
            logger.error(f"/1 持仓查询失败: {e}", exc_info=True)
            await update.message.reply_text(f"❌ 查询失败: {e}")

    @admin_only
    async def cmd_q_balances(update, ctx):
        """/2 全账户余额（合约+现货+资金，按 role 过滤；迁到 trader.multi 四账户）"""
        try:
            from shared.query_cards import render_all_wallets_card
            card = render_all_wallets_card(_ACC_ROLE)
            if len(card) > 4000:
                card = card[:4000] + "\n..."
            await update.message.reply_text(card, parse_mode=ParseMode.HTML)
        except PermissionError as e:
            await update.message.reply_text(f"⛔ {e}")
        except Exception as e:
            logger.error(f"/2 余额查询失败: {e}", exc_info=True)
            await update.message.reply_text(f"❌ 查询失败: {e}")

    async def _do_transfer(update, ctx, transfer_type, desc):
        import sys; sys.path.insert(0, f'/root/{bot_dir}')
        from trader.exchange import transfer_funds
        if not ctx.args:
            await update.message.reply_text(f"用法: {desc} 后跟金额，例如：/3 100")
            return
        try:
            amount = float(ctx.args[0])
        except ValueError:
            await update.message.reply_text("❌ 金额格式错误，请输入数字")
            return
        try:
            tran_id = transfer_funds(amount, transfer_type)
            await update.message.reply_text(f"✅ {desc} {amount} USDT\ntranId: {tran_id}")
        except Exception as e:
            await update.message.reply_text(f"❌ 划转失败: {e}")

    @admin_only
    async def cmd_transfer_3(update, ctx):
        await _do_transfer(update, ctx, "MAIN_UMFUTURE", "现货→合约")

    @admin_only
    async def cmd_transfer_4(update, ctx):
        await _do_transfer(update, ctx, "UMFUTURE_MAIN", "合约→现货")

    @admin_only
    async def cmd_transfer_5(update, ctx):
        await _do_transfer(update, ctx, "MAIN_FUNDING", "现货→资金")

    @admin_only
    async def cmd_transfer_6(update, ctx):
        await _do_transfer(update, ctx, "FUNDING_MAIN", "资金→现货")

    @admin_only
    async def cmd_trade_log(update, ctx):
        import sys; sys.path.insert(0, f'/root/{bot_dir}')
        from trader.trade_log import get_recent, format_for_tg
        n = 20
        if ctx.args:
            try: n = min(50, int(ctx.args[0]))
            except: pass
        entries = get_recent(n)
        await update.message.reply_text(format_for_tg(entries), parse_mode=ParseMode.HTML)

    @admin_only
    async def cmd_bot_log(update, ctx):
        import sys; sys.path.insert(0, f'/root/{bot_dir}')
        from trader.bot_log import get_recent_bot_events, format_bot_events_tg
        n = 20
        if ctx.args:
            try: n = min(50, int(ctx.args[0]))
            except: pass
        entries = get_recent_bot_events(n)
        await update.message.reply_text(format_bot_events_tg(entries), parse_mode=ParseMode.HTML)

    @admin_only
    async def cmd_sys_log(update, ctx):
        import sys; sys.path.insert(0, f'/root/{bot_dir}')
        from trader.bot_log import get_recent_sys_snapshots, format_sys_snapshot_tg
        n = 6
        if ctx.args:
            try: n = min(24, int(ctx.args[0]))
            except: pass
        entries = get_recent_sys_snapshots(n)
        await update.message.reply_text(format_sys_snapshot_tg(entries), parse_mode=ParseMode.HTML)

    def _bot_log(event, detail=""):
        if bot_dir in ("maomao", "tiantian"):
            try:
                import sys; sys.path.insert(0, f'/root/{bot_dir}')
                from trader.bot_log import log_bot_event
                log_bot_event(event, detail)
            except Exception:
                pass

    async def post_init(app):
        # 数字快捷键排最前
        num_cmds = []
        if bot_dir == "maomao":
            num_cmds = [
                BotCommand("1","全账户持仓"),
                BotCommand("2","全账户余额"),
                BotCommand("all","全账户净值汇总"),
                BotCommand("pos1","币安1 持仓"),
                BotCommand("bal1","币安1 余额"),
                BotCommand("pos2","币安2 持仓"),
                BotCommand("bal2","币安2 余额"),
                BotCommand("pos3","币安3 持仓（李红兵）"),
                BotCommand("bal3","币安3 余额（李红兵）"),
                BotCommand("pos4","币安4 持仓（专攻组六）"),
                BotCommand("bal4","币安4 余额（专攻组六）"),
                BotCommand("3","现货→合约 /3 <金额>"),
                BotCommand("4","合约→现货 /4 <金额>"),
                BotCommand("5","现货→资金 /5 <金额>"),
                BotCommand("6","资金→现货 /6 <金额>"),
                BotCommand("7","交易日志 /7 [条数]"),
                BotCommand("8","Bot运行事件 /8 [条数]"),
                BotCommand("9","系统快照 /9 [条数]"),
            ]
        elif bot_dir == "tiantian":
            num_cmds = [
                BotCommand("1","全账户持仓"),
                BotCommand("2","全账户余额"),
                BotCommand("pos2","币安2 持仓"),
                BotCommand("bal2","币安2 余额"),
                BotCommand("pos3","币安3 持仓（李红兵）"),
                BotCommand("bal3","币安3 余额（李红兵）"),
                BotCommand("pos4","币安4 持仓（专攻组六）"),
                BotCommand("bal4","币安4 余额（专攻组六）"),
                BotCommand("3","现货→合约 /3 <金额>"),
                BotCommand("4","合约→现货 /4 <金额>"),
                BotCommand("5","现货→资金 /5 <金额>"),
                BotCommand("6","资金→现货 /6 <金额>"),
            ]
        base_cmds = num_cmds + [
            BotCommand("start","状态"), BotCommand("help","帮助"),
            BotCommand("model","切换模型"),
            BotCommand("mode","查看状态"), BotCommand("status","四Bot服务状态"),
            BotCommand("ping","心跳"), BotCommand("log","查看日志"),
            BotCommand("restart","重启本Bot"), BotCommand("stop","关闭本Bot"),
        ]
        await app.bot.set_my_commands(base_cmds)
        state = load_state(bot_dir)
        logger.info(f"{bot_name} v4.3 | Claude Code代理 | {state['model']}")
        _bot_log("online", f"Claude Code代理 | {model_short(state['model'])}")
        try:
            await app.bot.send_message(chat_id=admin_id,
                text=f"✅ <b>{bot_name} 上线</b>\n\n模型: <code>{state['model']}</code>",
                parse_mode=ParseMode.HTML)
        except Exception as e:
            logger.warning(f"上线通知失败: {e}")

    logger.info(f"=== {bot_name} v4.3 启动 ===")
    app = Application.builder().token(bot_token).post_init(post_init).build()
    app.add_handler(CommandHandler("start", cmd_start))
    app.add_handler(CommandHandler("help", cmd_help))
    app.add_handler(CommandHandler("model", cmd_model))
    app.add_handler(CommandHandler("mode", cmd_mode))
    app.add_handler(CommandHandler("status", cmd_status))
    app.add_handler(CommandHandler("ping", cmd_ping))
    app.add_handler(CommandHandler("log", cmd_log))
    app.add_handler(CommandHandler("restart", cmd_restart))
    app.add_handler(CommandHandler("stop", cmd_stop_bot))
    if bot_dir in ("maomao", "tiantian"):
        app.add_handler(CommandHandler("1", cmd_q_positions))
        app.add_handler(CommandHandler("2", cmd_q_balances))
        app.add_handler(CommandHandler("3", cmd_transfer_3))
        app.add_handler(CommandHandler("4", cmd_transfer_4))
        app.add_handler(CommandHandler("5", cmd_transfer_5))
        app.add_handler(CommandHandler("6", cmd_transfer_6))
    # 多账户快捷卡片：/all /pos1-/pos4 /bal1-/bal4（按 role 过滤）
    if bot_dir in ("maomao", "tiantian"):
        app.add_handler(CommandHandler("all", cmd_all_card))
        app.add_handler(CommandHandler("pos1", cmd_pos1))
        app.add_handler(CommandHandler("pos2", cmd_pos2))
        app.add_handler(CommandHandler("pos3", cmd_pos3))
        app.add_handler(CommandHandler("pos4", cmd_pos4))
        app.add_handler(CommandHandler("bal1", cmd_bal1))
        app.add_handler(CommandHandler("bal2", cmd_bal2))
        app.add_handler(CommandHandler("bal3", cmd_bal3))
        app.add_handler(CommandHandler("bal4", cmd_bal4))
    if bot_dir == "maomao":
        app.add_handler(CommandHandler("7", cmd_trade_log))
        app.add_handler(CommandHandler("8", cmd_bot_log))
        app.add_handler(CommandHandler("9", cmd_sys_log))
    from shared.message_buffer import make_buffered_handler
    buffered_handle_text = make_buffered_handler(handle_text, admin_id=admin_id)
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND & filters.ChatType.PRIVATE, buffered_handle_text))
    app.add_handler(MessageHandler(filters.PHOTO & filters.ChatType.PRIVATE, handle_photo))
    app.add_handler(MessageHandler(filters.Document.ALL & filters.ChatType.PRIVATE, handle_document))
    # 2026-04-27 视频/音频
    app.add_handler(MessageHandler((filters.VIDEO | filters.VIDEO_NOTE) & filters.ChatType.PRIVATE, handle_video))
    app.add_handler(MessageHandler((filters.VOICE | filters.AUDIO) & filters.ChatType.PRIVATE, handle_audio))
    from telegram.ext import CallbackQueryHandler
    app.add_handler(CallbackQueryHandler(handle_callback))
    app.add_error_handler(error_handler)

    async def _heartbeat(ctx):
        logger.info(f"[heartbeat] {bot_name} alive")
        if bot_dir in ("maomao", "tiantian"):
            try:
                import sys; sys.path.insert(0, f'/root/{bot_dir}')
                from trader.bot_log import log_sys_snapshot
                log_sys_snapshot()
            except Exception:
                pass
    app.job_queue.run_repeating(_heartbeat, interval=300, first=10)

    try:
        app.run_polling(drop_pending_updates=True)
    finally:
        _bot_log("offline", "正常退出")
        import httpx
        try:
            httpx.post(f"https://api.telegram.org/bot{bot_token}/sendMessage",
                json={"chat_id": admin_id, "text": f"🔴 <b>{bot_name} 下线</b>", "parse_mode": "HTML"}, timeout=5)
        except Exception:
            pass
